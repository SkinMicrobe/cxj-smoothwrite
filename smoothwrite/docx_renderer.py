"""Render smoothwrite Markdown articles and figure plans into Word documents."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from .io_utils import ensure_dir, load_default_config, read_json, read_text
from .models import FigurePlan, FigurePlanItem

_SECTION_PATTERN = re.compile(
    r"^(?:#{1,6}\s*)?([一二三四五六七八九十百]+、.+)$"
)
_NUMBERED_POINT_PATTERN = re.compile(r"^\s*\d+[、）)].+")
_FIGURE_MARKER_PATTERN = re.compile(
    r"^\s*(配图\d+)\s*｜\s*((?:Figure|Fig\.?)\s*\d+[A-Za-z]?(?:\s*[–—-]\s*[A-Za-z])?)\s*$",
    flags=re.IGNORECASE,
)
_SMALL_TEXT_PATTERN = re.compile(r"^(论文信息|原文链接)：")
_INLINE_MARKDOWN_PATTERN = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def _set_run_font(run: object, size: float, bold: bool = False) -> None:
    """Apply a readable Chinese-capable font to a run."""
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run._element.rPr.rFonts.set(qn("w:cs"), "微软雅黑")
    run._element.rPr.rFonts.set(qn("w:hint"), "eastAsia")
    run.font.size = Pt(size)
    size_cs = OxmlElement("w:szCs")
    size_cs.set(qn("w:val"), str(int(size * 2)))
    run._element.rPr.append(size_cs)
    run.bold = bold


def _configure_document_styles(document: Document) -> None:
    """Set a restrained reference-matched style for WeChat editorial drafts."""
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(10)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title = document.styles["Title"]
    title.font.name = "微软雅黑"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    title.font.size = Pt(15)
    title.font.bold = True

    heading = document.styles["Heading 1"]
    heading.font.name = "微软雅黑"
    heading._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.paragraph_format.space_before = Pt(14)
    heading.paragraph_format.space_after = Pt(7)


def _add_title(document: Document, text: str) -> None:
    """Add the article title."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run(text)
    _set_run_font(run, size=15, bold=True)


def _add_heading(document: Document, text: str) -> None:
    """Add a Chinese-numbered section heading."""
    paragraph = document.add_paragraph(style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    _set_run_font(run, size=12, bold=True)


def _add_markdown_runs(paragraph: object, text: str, size: float) -> None:
    """Render the small inline Markdown subset used by smoothwrite."""
    cursor = 0
    for match in _INLINE_MARKDOWN_PATTERN.finditer(text):
        if match.start() > cursor:
            _set_run_font(paragraph.add_run(text[cursor : match.start()]), size=size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            _set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            _set_run_font(run, size=size)
            run.italic = True
        cursor = match.end()
    if cursor < len(text):
        _set_run_font(paragraph.add_run(text[cursor:]), size=size)


def _configure_cjk_paragraph(paragraph: object) -> None:
    """Add the East Asian line-layout flags used by the reference DOCX."""
    properties = paragraph._p.get_or_add_pPr()
    for tag in (
        "w:kinsoku",
        "w:wordWrap",
        "w:overflowPunct",
        "w:autoSpaceDE",
        "w:autoSpaceDN",
        "w:adjustRightInd",
        "w:snapToGrid",
    ):
        if properties.find(qn(tag)) is None:
            properties.append(OxmlElement(tag))


def _add_body_paragraph(
    document: Document,
    text: str,
    *,
    indent: bool = False,
    small: bool = False,
) -> None:
    """Add a normal article paragraph."""
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _configure_cjk_paragraph(paragraph)
    if indent:
        paragraph.paragraph_format.left_indent = Cm(0.5)
    _add_markdown_runs(paragraph, text, size=9.5 if small else 12)


def _figure_lookup(figure_plan: FigurePlan) -> dict[str, FigurePlanItem]:
    """Build label and source-figure lookups for planned figures."""
    lookup: dict[str, FigurePlanItem] = {}
    for figure in figure_plan.figures:
        lookup[figure.label.casefold()] = figure
        lookup[figure.source_figure.casefold()] = figure
    return lookup


def _resolve_image_path(
    image_path: str | None,
    base_path: Path,
) -> Path | None:
    """Resolve a figure image path relative to the selected base directory."""
    if not image_path:
        return None
    path = Path(image_path)
    return path if path.is_absolute() else base_path / path


def _add_figure(
    document: Document,
    marker_text: str,
    label: str,
    source_figure: str,
    figures: dict[str, FigurePlanItem],
    base_path: Path,
    image_width_cm: float,
) -> None:
    """Add an image followed by a source marker and explanatory caption."""
    figure = figures.get(label.casefold()) or figures.get(source_figure.casefold())
    image_path = _resolve_image_path(figure.image_path if figure else None, base_path)

    if image_path is not None and image_path.is_file():
        image_paragraph = document.add_paragraph()
        image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = image_paragraph.add_run()
        try:
            run.add_picture(str(image_path), width=Cm(image_width_cm))
        except Exception:
            image_path = None
    if image_path is None or not image_path.is_file():
        fallback = document.add_paragraph()
        fallback.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = fallback.add_run("[图片未找到，请人工插入]")
        _set_run_font(run, size=10)

    marker = document.add_paragraph()
    marker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    marker.paragraph_format.space_after = Pt(2)
    marker_run = marker.add_run(marker_text)
    _set_run_font(marker_run, size=9.5, bold=True)
    if figure and figure.caption:
        caption = document.add_paragraph()
        caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
        caption.paragraph_format.space_after = Pt(8)
        caption_run = caption.add_run(figure.caption)
        _set_run_font(caption_run, size=9)
        caption_run.font.color.rgb = RGBColor(80, 80, 80)


def render_docx(
    article_md_path: str,
    figure_plan_path: str,
    output_docx_path: str,
    base_dir: str | None = None,
) -> str:
    """Render article Markdown and a FigurePlan JSON file into a DOCX."""
    article_path = Path(article_md_path)
    plan_path = Path(figure_plan_path)
    output_path = Path(output_docx_path)

    if not article_path.is_file():
        raise FileNotFoundError(f"Article Markdown file does not exist: {article_path}")
    if not plan_path.is_file():
        raise FileNotFoundError(f"Figure plan file does not exist: {plan_path}")

    article = read_text(article_path)
    figure_plan = FigurePlan.model_validate(read_json(plan_path))
    figures = _figure_lookup(figure_plan)
    base_path = Path(base_dir) if base_dir is not None else plan_path.parent
    image_width_cm = float(load_default_config()["figures"]["image_width_cm"])

    document = Document()
    _configure_document_styles(document)

    lines = article.splitlines()
    first_content_written = False

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        if not first_content_written and stripped.startswith("# "):
            _add_title(document, stripped[2:].strip())
            first_content_written = True
            continue
        first_content_written = True

        section_match = _SECTION_PATTERN.match(stripped)
        if section_match:
            _add_heading(document, section_match.group(1))
            continue

        figure_match = _FIGURE_MARKER_PATTERN.match(stripped)
        if figure_match:
            _add_figure(
                document=document,
                marker_text=stripped,
                label=figure_match.group(1),
                source_figure=figure_match.group(2),
                figures=figures,
                base_path=base_path,
                image_width_cm=image_width_cm,
            )
            continue

        _add_body_paragraph(
            document,
            stripped,
            indent=bool(_NUMBERED_POINT_PATTERN.match(stripped)),
            small=bool(_SMALL_TEXT_PATTERN.match(stripped)),
        )

    ensure_dir(output_path.parent)
    document.save(str(output_path))
    return str(output_path)


def render_docx_from_output_dir(output_dir: str) -> str:
    """Render the conventional smoothwrite output directory into final.docx."""
    output_path = Path(output_dir)
    output_config = load_default_config()["output"]
    return render_docx(
        article_md_path=str(output_path / output_config["article_md"]),
        figure_plan_path=str(output_path / output_config["figure_plan_json"]),
        output_docx_path=str(output_path / output_config["final_docx"]),
        base_dir=str(output_path),
    )
